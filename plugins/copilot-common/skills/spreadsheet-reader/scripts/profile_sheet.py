#!/usr/bin/env python3
"""Profile a CSV/XLSX file locally: columns, types, fill rates, stats, key candidates.

Usage: profile_sheet.py FILE [--sheet NAME] [--rows N] [--json]
Stdlib only; .xlsx needs openpyxl (optional). No network access.
"""
import argparse
import csv
import json
import os
import re
import statistics
import sys
from collections import Counter

NUM_RE = re.compile(r"^-?\$?\s*[\d,]*\.?\d+\s*%?$")
DATE_RE = re.compile(r"^(\d{1,4}[-/]\d{1,2}[-/]\d{1,4})(\s+\d{1,2}:\d{2}(:\d{2})?)?$")


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not NUM_RE.match(s):
        return None
    s = s.replace("$", "").replace(",", "").replace("%", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def load_csv(path):
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(f, dialect))
    return {os.path.basename(path): rows}


def load_xlsx(path, sheet=None):
    try:
        import openpyxl  # type: ignore
    except ImportError:
        sys.exit("openpyxl not installed. Either `pip install openpyxl` or export the sheet to CSV.")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = [sheet] if sheet else wb.sheetnames
    out = {}
    for n in names:
        ws = wb[n]
        out[n] = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
    return out


def profile(rows, sample_rows):
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return {"rows": 0, "columns": []}
    header = [str(h).strip() or f"col{i+1}" for i, h in enumerate(rows[0])]
    body = rows[1:]
    width = len(header)
    cols = []
    for i, name in enumerate(header):
        vals = [r[i] if i < len(r) else "" for r in body]
        filled = [v for v in vals if str(v).strip() != ""]
        nums = [to_num(v) for v in filled]
        nums = [n for n in nums if n is not None]
        dates = sum(1 for v in filled if DATE_RE.match(str(v).strip()) or type(v).__name__ in ("datetime", "date"))
        n = len(filled)
        if n and len(nums) / n > 0.9:
            kind = "number"
        elif n and dates / n > 0.9:
            kind = "date"
        else:
            kind = "text"
        distinct = len(set(str(v).strip().lower() for v in filled))
        col = {
            "name": name,
            "type": kind,
            "fill_rate": round(n / len(body), 3) if body else 0,
            "distinct": distinct,
            "key_candidate": kind == "text" and bool(body) and n == len(body) and distinct == n and n > 1,
        }
        if kind == "number" and nums:
            col.update(
                min=min(nums), max=max(nums), sum=round(sum(nums), 2),
                mean=round(statistics.fmean(nums), 4),
                negatives=sum(1 for x in nums if x < 0), zeros=sum(1 for x in nums if x == 0),
            )
        else:
            top = Counter(str(v).strip() for v in filled).most_common(5)
            col["top_values"] = top
        cols.append(col)
    return {
        "rows": len(body),
        "width": width,
        "columns": cols,
        "sample": [[str(c) for c in r[:width]] for r in body[:sample_rows]],
        "header": header,
    }


def render(name, p):
    print(f"=== SHEET: {name}  rows={p['rows']} cols={p.get('width', 0)}")
    if not p["rows"]:
        print("  (empty)")
        return
    for c in p["columns"]:
        line = f"  - {c['name']:<28} {c['type']:<7} fill={c['fill_rate']:<6} distinct={c['distinct']}"
        if c["key_candidate"]:
            line += "  [KEY?]"
        if c["type"] == "number" and "sum" in c:
            line += f"  min={c['min']} max={c['max']} sum={c['sum']} mean={c['mean']}"
            if c["negatives"]:
                line += f" negatives={c['negatives']}"
            if c["zeros"]:
                line += f" zeros={c['zeros']}"
        else:
            tv = ", ".join(f"{v!r}×{k}" for v, k in c.get("top_values", [])[:3])
            line += f"  top: {tv}"
        print(line)
    print("  SAMPLE:")
    print("   | " + " | ".join(p["header"]))
    for r in p["sample"]:
        print("   | " + " | ".join(r))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--sheet")
    ap.add_argument("--rows", type=int, default=5)
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    ext = os.path.splitext(a.file)[1].lower()
    if ext in (".csv", ".txt", ".tsv"):
        sheets = load_csv(a.file)
    elif ext in (".xlsx", ".xlsm"):
        sheets = load_xlsx(a.file, a.sheet)
    elif ext == ".xls":
        sys.exit("Legacy .xls is not supported locally; save as .xlsx or .csv first.")
    else:
        sys.exit(f"Unsupported extension {ext}")
    result = {n: profile(r, a.rows) for n, r in sheets.items()}
    if a.json:
        print(json.dumps(result, default=str, indent=2))
    else:
        for n, p in result.items():
            render(n, p)


if __name__ == "__main__":
    main()
